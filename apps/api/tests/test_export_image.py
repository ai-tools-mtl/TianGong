"""导出 image 节点渲染测试。"""


def test_tiptap_to_markdown_renders_image():
    from app.services.export_service import _tiptap_to_markdown

    doc_json = {"type": "doc", "content": [
        {"type": "image", "attrs": {"src": "/api/v1/x/file", "alt": "图 1 示意图"}}
    ]}
    md = _tiptap_to_markdown(doc_json)
    assert "![图 1 示意图](/api/v1/x/file)" in md


def test_render_tiptap_to_docx_handles_image(db_session, monkeypatch):
    """图片节点:从附件 URL 反查 → minio 取字节 → add_picture(用 BytesIO)。"""
    import uuid as _uuid

    from docx import Document
    from sqlalchemy import select

    from app.core.security import hash_password
    from app.models import Attachment, Project, User
    from app.services import export_service

    # python-docx 的 add_picture 在 docx.document.Document 类上
    from docx.document import Document as DocClass

    # 造 user + project + attachment,storage_path 指向 mock 的 key
    u = User(email="img@example.com", password_hash=hash_password("P1!"), name="I")
    db_session.add(u)
    db_session.flush()
    p = Project(user_id=u.id, title="t")
    db_session.add(p)
    db_session.flush()
    att = Attachment(
        project_id=p.id, filename="f.png",
        storage_path="attachments/x/y.png",
        mime_type="image/png", size=8,
    )
    db_session.add(att)
    db_session.commit()

    # mock storage 返回 PNG 头(足够让 add_picture 被调用,内容无所谓)
    from app.core.storage import get_storage

    storage = get_storage()
    storage.put("personal", "attachments/x/y.png", b"\x89PNG\r\n\x1a\n", "image/png")

    called = {"add_picture": False}

    def fake_add_picture(self, *a, **kw):
        called["add_picture"] = True

    monkeypatch.setattr(DocClass, "add_picture", fake_add_picture)

    doc = Document()
    # src 形如真实下载 URL
    src = f"/api/v1/projects/{p.id}/attachments/{att.id}/file"
    doc_json = {"type": "doc", "content": [
        {"type": "image", "attrs": {"src": src, "alt": "图1"}}
    ]}
    export_service._render_tiptap_to_docx(doc, doc_json, db_session)
    assert called["add_picture"] is True


def test_render_tiptap_to_docx_image_missing_attachment(db_session):
    """附件不存在时,图片分支静默跳过,不阻断导出(关键约束:try/except 兜底)。"""
    from docx import Document

    from app.services import export_service

    doc = Document()
    src = "/api/v1/projects/00000000-0000-0000-0000-000000000000/attachments/00000000-0000-0000-0000-000000000000/file"
    doc_json = {"type": "doc", "content": [
        {"type": "image", "attrs": {"src": src, "alt": "缺失图"}}
    ]}
    # 不抛即通过
    export_service._render_tiptap_to_docx(doc, doc_json, db_session)
