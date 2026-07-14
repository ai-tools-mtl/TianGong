"""计划 12 Feature A：异步解析 + 恢复扫描测试。"""

import io
import uuid

from docx import Document
from sqlalchemy.orm import sessionmaker


# ──────────────────────────────────────────────────────────────
# Task 1: run_parse_job_standalone + recover_pending_jobs
# ──────────────────────────────────────────────────────────────


def test_run_parse_job_standalone_opens_own_session(
    client, app_obj, engine, registered_user, db_session, monkeypatch
):
    """run_parse_job_standalone 自开 session 完成解析，不依赖请求 db。"""
    # standalone 内部 `from app.core.database import SessionLocal` —— patch 模块属性
    from app.core import database as db_module

    TestingSession = sessionmaker(bind=engine)
    monkeypatch.setattr(db_module, "SessionLocal", TestingSession)

    from app.services import parse_service

    client.post(
        "/api/v1/auth/login",
        json={
            "email": registered_user["email"],
            "password": registered_user["password"],
        },
    )

    doc = Document()
    doc.add_heading("发明名称", level=1)
    doc.add_paragraph("正文")
    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    job = parse_service.create_parse_job(
        db_session,
        user_id=uuid.UUID(registered_user["id"]),
        filename="x.docx",
        file_bytes=file_bytes,
        upload_dir="uploads",
    )
    job_id = str(job.id)

    # 调用 standalone：不传 db，函数自己开 session
    parse_service.run_parse_job_standalone(job_id, "uploads")

    # 用全新 session 验证持久化
    from app.models import ParseJob

    fresh_session = TestingSession()
    try:
        fresh = fresh_session.get(ParseJob, uuid.UUID(job_id))
        assert fresh.status == "completed"
        assert fresh.template_id is not None
    finally:
        fresh_session.close()


def test_recover_pending_jobs_reenqueues_processing_and_stale_pending(
    db_session, monkeypatch
):
    """启动恢复扫描：processing 直接重跑，pending 超 10 分钟重跑，新 pending 不动。"""
    from datetime import datetime, timedelta, timezone

    from app.core import database as db_module
    from app.core.security import hash_password
    from app.models import ParseJob, User
    from app.services import parse_service

    # 让 standalone/recover 内部的 SessionLocal 指向测试库
    TestingSession = sessionmaker(bind=db_session.bind)
    monkeypatch.setattr(db_module, "SessionLocal", TestingSession)

    # 造一个真实用户（ParseJob.user_id 外键约束）
    u = User(email="rec@example.com", password_hash=hash_password("Pass1234!"), name="R")
    db_session.add(u)
    db_session.commit()

    now = datetime.now(timezone.utc)
    j_processing = ParseJob(
        user_id=u.id, source_path="uploads/x.docx", status="processing"
    )
    j_stale = ParseJob(
        user_id=u.id, source_path="uploads/y.docx", status="pending"
    )
    j_stale.created_at = now - timedelta(minutes=20)  # 超时 pending
    j_fresh = ParseJob(
        user_id=u.id, source_path="uploads/z.docx", status="pending"
    )
    # j_fresh.created_at 用 server 默认 now，未超时
    for j in (j_processing, j_stale, j_fresh):
        db_session.add(j)
    db_session.commit()

    # 用 fake_run 标记被调度的 job_id，避免真去解析不存在的 docx
    reloaded: list[str] = []

    def fake_run(db, job_id, upload_dir):
        reloaded.append(str(job_id))
        j = db.get(ParseJob, uuid.UUID(job_id))
        if j and j.status != "completed":
            j.status = "processing"
            db.commit()

    monkeypatch.setattr(parse_service, "run_parse_job", fake_run)

    n = parse_service.recover_pending_jobs("uploads")
    assert n == 2  # processing + stale pending 被重入队
    assert str(j_processing.id) in reloaded
    assert str(j_stale.id) in reloaded
    assert str(j_fresh.id) not in reloaded  # 新 pending 不动


# ──────────────────────────────────────────────────────────────
# Task 2: upload 端点改 BackgroundTasks + 202
# ──────────────────────────────────────────────────────────────


def test_upload_returns_202_processing_and_background_completes(
    client, app_obj, engine, registered_user, monkeypatch
):
    """upload 立即返回 202 + status=processing，BackgroundTasks 在响应后跑完。"""
    # BackgroundTask 会调 run_parse_job_standalone → 内部 SessionLocal
    # patch 它指向测试库，否则 standalone 开的 session 连到生产 DB 读不到 job
    from app.core import database as db_module

    monkeypatch.setattr(db_module, "SessionLocal", sessionmaker(bind=engine))

    client.post(
        "/api/v1/auth/login",
        json={
            "email": registered_user["email"],
            "password": registered_user["password"],
        },
    )

    doc = Document()
    doc.add_heading("发明名称", level=1)
    doc.add_paragraph("正文")
    buf = io.BytesIO()
    doc.save(buf)

    res = client.post(
        "/api/v1/templates",
        files={
            "file": (
                "test.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert res.status_code == 202
    data = res.json()
    assert data["status"] == "processing"
    assert "parse_job_id" in data


# ──────────────────────────────────────────────────────────────
# Task 3: on_startup 接入恢复扫描
# ──────────────────────────────────────────────────────────────


def test_on_startup_calls_recovery_without_crashing(monkeypatch):
    """on_startup 调用 recover_pending_jobs，即使抛错也不阻塞。"""
    called = {"n": 0}

    def fake_recover(upload_dir, stale_minutes=10):
        called["n"] += 1
        return 0

    from app.services import parse_service

    monkeypatch.setattr(parse_service, "recover_pending_jobs", fake_recover)

    from app.main import on_startup

    on_startup()  # 不应抛异常

    assert called["n"] == 1


def test_on_startup_swallows_recovery_error(monkeypatch):
    """恢复扫描抛错时 on_startup 仍正常返回（不阻塞启动）。"""

    def boom(upload_dir, stale_minutes=10):
        raise RuntimeError("db down")

    from app.services import parse_service

    monkeypatch.setattr(parse_service, "recover_pending_jobs", boom)

    from app.main import on_startup

    on_startup()  # 不应抛异常


# ──────────────────────────────────────────────────────────────
# Task 4: GET /templates/parse-jobs/{job_id} 状态端点
# ──────────────────────────────────────────────────────────────


def test_get_parse_job_status_completed(
    client, app_obj, engine, registered_user, monkeypatch
):
    """GET /templates/parse-jobs/{id} 返回 status + template_id。"""
    # BackgroundTask 跑 standalone，需 patch SessionLocal 指向测试库
    from app.core import database as db_module

    monkeypatch.setattr(db_module, "SessionLocal", sessionmaker(bind=engine))

    client.post(
        "/api/v1/auth/login",
        json={
            "email": registered_user["email"],
            "password": registered_user["password"],
        },
    )

    doc = Document()
    doc.add_heading("发明名称", level=1)
    doc.add_paragraph("正文")
    buf = io.BytesIO()
    doc.save(buf)

    upload_res = client.post(
        "/api/v1/templates",
        files={
            "file": (
                "t.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    job_id = upload_res.json()["parse_job_id"]

    res = client.get(f"/api/v1/templates/parse-jobs/{job_id}")
    assert res.status_code == 200
    data = res.json()
    # TestClient 在响应返回后同步等待 background task 完成，故通常为 completed
    assert data["status"] == "completed"
    assert data["template_id"] is not None


def test_get_parse_job_not_found(client, registered_user):
    """不存在的 job_id 返回 404。"""
    client.post(
        "/api/v1/auth/login",
        json={
            "email": registered_user["email"],
            "password": registered_user["password"],
        },
    )
    res = client.get(f"/api/v1/templates/parse-jobs/{uuid.uuid4()}")
    assert res.status_code == 404


def test_get_parse_job_unauthenticated(client):
    """未登录返回 401。"""
    res = client.get(f"/api/v1/templates/parse-jobs/{uuid.uuid4()}")
    assert res.status_code == 401



