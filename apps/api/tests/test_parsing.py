"""Word 解析器单元测试。"""

from docx import Document

from app.parsing.structure_extractor import extract_structure


def _make_docx_with_headings():
    """构造一个含 Heading 样式段落的 docx（内存）。"""
    doc = Document()
    doc.add_heading("发明名称", level=1)
    doc.add_paragraph("一些正文内容")
    doc.add_heading("技术领域", level=1)
    doc.add_heading("现有技术", level=2)
    doc.add_paragraph("更多正文")
    return doc


def test_extract_structure_from_headings():
    doc = _make_docx_with_headings()
    sections = extract_structure(doc)
    assert len(sections) == 3
    assert sections[0]["title"] == "发明名称"
    assert sections[0]["level"] == 1
    assert sections[2]["title"] == "现有技术"
    assert sections[2]["level"] == 2


def test_extract_structure_assigns_keys_by_title():
    """标题匹配已知章节时自动分配 key。"""
    doc = _make_docx_with_headings()
    sections = extract_structure(doc)
    assert sections[0]["key"] == "name"
    assert sections[1]["key"] == "field"


def test_extract_structure_unknown_title_gets_custom_key():
    doc = Document()
    doc.add_heading("某自定义章节", level=1)
    sections = extract_structure(doc)
    assert sections[0]["key"] == "custom"


def test_extract_structure_no_headings_falls_back_to_empty():
    """无 Heading 时返回空列表。"""
    doc = Document()
    doc.add_paragraph("只有正文，没有标题")
    sections = extract_structure(doc)
    assert sections == []


# ── 样式提取 ──

def test_extract_styles_basic():
    from app.parsing.style_extractor import extract_styles
    doc = _make_docx_with_headings()
    styles = extract_styles(doc)
    assert isinstance(styles, dict)
    assert len(styles) > 0
    for key, val in styles.items():
        assert isinstance(val, dict)


# ── 编号解析 ──

def test_resolve_numbering_from_text_fallback():
    """正则兜底：段落文本以数字开头时提取编号。"""
    from app.parsing.numbering_resolver import extract_numbering_from_text
    assert extract_numbering_from_text("1.1 这是章节") == "1.1"
    assert extract_numbering_from_text("1.1.2 内容") == "1.1.2"
    assert extract_numbering_from_text("2、标题") == "2"
    assert extract_numbering_from_text("无编号的正文") is None


def test_resolve_numbering_no_element_returns_none():
    """无 numbering.xml 时优雅降级，不报错。"""
    from app.parsing.numbering_resolver import NumberingResolver
    resolver = NumberingResolver(None)
    assert resolver.resolve_for_paragraph(None) is None


# ── 解析编排集成 ──

def test_parse_docx_integration():
    """端到端：构造 docx → 解析 → 得到 structure/styles/numbering。"""
    from app.parsing.docx_parser import parse_docx
    doc = _make_docx_with_headings()
    parsed = parse_docx(doc)
    assert len(parsed.structure) == 3
    assert parsed.structure[0]["key"] == "name"
    assert isinstance(parsed.styles, dict)
    assert isinstance(parsed.numbering, dict)
