"""导出服务：Tiptap JSON → docx / Markdown（设计 13.5）。

存储改造(T3):图片从本地路径改为从 minio 取字节流。
"""

import io
import re
import uuid
from typing import Any

from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Attachment, Project, Section


def export_markdown(db: Session, *, project: Project) -> str:
    """导出为 Markdown 文本。"""
    sections = _get_ordered_sections(db, project)
    lines = [f"# {project.title}\n"]

    if project.metadata_:
        meta = project.metadata_
        if meta.get("inventors"):
            lines.append(f"**发明人**：{', '.join(meta['inventors'])}\n")
        if meta.get("applicant"):
            lines.append(f"**申请人**：{meta['applicant']}\n")

    for s in sections:
        lines.append(f"\n## {s.title}\n")
        if s.content:
            lines.append(_tiptap_to_markdown(db, s.content))
        lines.append("")
    return "\n".join(lines)


def export_docx(db: Session, *, project: Project) -> bytes:
    """导出为 .docx 字节。"""
    doc = Document()
    doc.add_heading(project.title, level=0)

    if project.metadata_:
        meta = project.metadata_
        info_parts = []
        if meta.get("inventors"):
            info_parts.append(f"发明人：{', '.join(meta['inventors'])}")
        if meta.get("applicant"):
            info_parts.append(f"申请人：{meta['applicant']}")
        if meta.get("disclosure_date"):
            info_parts.append(f"日期：{meta['disclosure_date']}")
        if info_parts:
            doc.add_paragraph(" | ".join(info_parts))

    sections = _get_ordered_sections(db, project)
    for s in sections:
        doc.add_heading(s.title, level=1)
        if s.content:
            _render_tiptap_to_docx(doc, s.content, db)
        else:
            doc.add_paragraph("（待填写）")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _get_ordered_sections(db: Session, project: Project) -> list[Section]:
    return list(db.scalars(
        select(Section).where(Section.project_id == project.id).order_by(Section.order)
    ))


def _tiptap_to_markdown(db: Session, doc_json: dict) -> str:
    """Tiptap JSON → Markdown 纯文本。"""
    parts: list[str] = []

    def walk(node: Any):
        if isinstance(node, dict):
            ntype = node.get("type")
            if ntype == "text":
                text = node.get("text", "")
                marks = node.get("marks", [])
                if any(m.get("type") == "bold" for m in marks):
                    text = f"**{text}**"
                parts.append(text)
            elif ntype == "paragraph":
                for child in node.get("content", []):
                    walk(child)
                parts.append("\n\n")
            elif ntype == "heading":
                level = node.get("attrs", {}).get("level", 2)
                parts.append(f"\n{'#' * level} ")
                for child in node.get("content", []):
                    walk(child)
                parts.append("\n\n")
            elif ntype == "image":
                src = node.get("attrs", {}).get("src", "")
                alt = node.get("attrs", {}).get("alt", "")
                # 鉴权 URL 无法在外部 Markdown 阅读器显示，转 base64 data URI 自包含
                img_bytes = _fetch_image_bytes(db, src)
                if img_bytes:
                    import base64
                    b64 = base64.b64encode(img_bytes).decode("ascii")
                    mime = "image/png"
                    if ".jp" in src:
                        mime = "image/jpeg"
                    elif ".gif" in src:
                        mime = "image/gif"
                    parts.append(f"\n\n![{alt}](data:{mime};base64,{b64})\n\n")
                else:
                    parts.append(f"\n\n![{alt}]({src})\n\n")
            elif ntype in ("bulletList", "orderedList"):
                for child in node.get("content", []):
                    walk(child)
            elif ntype == "listItem":
                parts.append("- ")
                for child in node.get("content", []):
                    walk(child)
                parts.append("\n")
            elif ntype == "table":
                # 收集所有行数据再统一输出为 GFM 表格
                rows_data = _collect_table_rows(node)
                if rows_data:
                    col_count = max(len(r) for r in rows_data)
                    parts.append("\n")
                    for row_idx, row in enumerate(rows_data):
                        padded = row + [""] * (col_count - len(row))
                        parts.append("| " + " | ".join(padded) + " |\n")
                        # 首行（表头）后插入分隔行
                        if row_idx == 0:
                            parts.append("|" + "|".join(["---"] * col_count) + "|\n")
                    parts.append("\n")
            elif ntype in ("tableHeader", "tableCell"):
                for child in node.get("content", []):
                    walk(child)
            elif ntype == "tableRow":
                # 由 table 节点统一处理，此处跳过以避免重复输出
                pass
            else:
                for child in node.get("content", []):
                    walk(child)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(doc_json)
    return "".join(parts).strip()


def _collect_table_rows(table_node: dict) -> list[list[str]]:
    """从 Tiptap table 节点提取二维文本数组，供 _tiptap_to_markdown 使用。"""
    rows: list[list[str]] = []
    for row in table_node.get("content", []):
        if row.get("type") != "tableRow":
            continue
        cells: list[str] = []
        for cell in row.get("content", []):
            cell_text = _extract_node_text(cell)
            cells.append(cell_text)
        rows.append(cells)
    return rows


def _extract_node_text(node: dict) -> str:
    """递归提取节点内所有 text 节点的文本（保留加粗 ** 标记）。"""
    parts: list[str] = []

    def _walk(n: Any) -> None:
        if isinstance(n, dict):
            if n.get("type") == "text":
                t = n.get("text", "")
                if any(m.get("type") == "bold" for m in n.get("marks", [])):
                    t = f"**{t}**"
                parts.append(t)
            else:
                for child in n.get("content", []):
                    _walk(child)
        elif isinstance(n, list):
            for item in n:
                _walk(item)

    _walk(node)
    return "".join(parts)


def _render_tiptap_to_docx(doc: Document, doc_json: dict, db: Session) -> None:
    """把 Tiptap JSON 渲染到 python-docx Document。

    图片节点:src 形如 /api/v1/projects/{pid}/attachments/{aid}/file,
    从中提 attachment_id → 查 storage_path → 从 minio 取字节 → add_picture。
    图片缺失不阻断导出(try/except 兜底)。
    """
    def walk(node: Any):
        if isinstance(node, dict):
            ntype = node.get("type")
            if ntype == "paragraph":
                texts = [_get_text(c) for c in node.get("content", [])]
                doc.add_paragraph("".join(texts))
            elif ntype == "heading":
                level = node.get("attrs", {}).get("level", 2)
                texts = [_get_text(c) for c in node.get("content", [])]
                doc.add_heading("".join(texts), level=min(level, 3))
            elif ntype == "image":
                src = node.get("attrs", {}).get("src", "")
                alt = node.get("attrs", {}).get("alt", "")
                buf = _fetch_image_bytes(db, src)
                if buf:
                    try:
                        doc.add_picture(io.BytesIO(buf))
                    except Exception:
                        pass  # 图片损坏不阻断导出
                doc.add_paragraph(alt)
            elif ntype in ("bulletList", "orderedList"):
                style = "List Bullet" if ntype == "bulletList" else "List Number"
                for item in node.get("content", []):
                    if item.get("type") == "listItem":
                        texts = [_get_text(c) for c in item.get("content", [])]
                        doc.add_paragraph("".join(texts), style=style)
            elif ntype == "table":
                rows_data = _collect_table_rows(node)
                if rows_data:
                    col_count = max(len(r) for r in rows_data)
                    table = doc.add_table(rows=len(rows_data), cols=col_count)
                    table.style = "Table Grid"
                    for i, row_data in enumerate(rows_data):
                        for j, cell_text in enumerate(row_data):
                            if j < col_count:
                                cell = table.cell(i, j)
                                cell.text = cell_text
                                # 首行（表头）加粗
                                if i == 0:
                                    for paragraph in cell.paragraphs:
                                        for run in paragraph.runs:
                                            run.bold = True
            else:
                for child in node.get("content", []):
                    walk(child)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(doc_json)


# 匹配附件下载 URL 里的 attachment_id(UUID)
_ATT_URL_RE = re.compile(r"/attachments/([0-9a-fA-F-]{36})/file")


def _fetch_image_bytes(db: Session, src: str) -> bytes | None:
    """从 Tiptap image src 反查 attachment,从 minio 取图片字节。失败返回 None。"""
    if not src:
        return None
    m = _ATT_URL_RE.search(src)
    if not m:
        return None
    try:
        att_id = uuid.UUID(m.group(1))
    except ValueError:
        return None
    att = db.get(Attachment, att_id)
    if att is None or not att.storage_path:
        return None
    try:
        from app.core.storage import get_storage

        return get_storage().get("personal", att.storage_path)
    except Exception:
        return None


def fetch_image_bytes_for_project(
    db: Session, src: str, project_id: uuid.UUID
) -> tuple[bytes, str] | None:
    """从 Tiptap image src 反查 attachment，校验项目归属后返回 (字节, mime)。

    与 _fetch_image_bytes 的区别：增加 att.project_id == project_id 归属校验，
    用于游客浏览等公开场景，防止持有某项目 token 的访客枚举其他项目的附图。
    跨项目 / 外部 URL / 取不到字节 → 一律 None（调用方按需置空，不泄露）。
    """
    if not src:
        return None
    m = _ATT_URL_RE.search(src)
    if not m:
        return None
    try:
        att_id = uuid.UUID(m.group(1))
    except ValueError:
        return None
    att = db.get(Attachment, att_id)
    if att is None or not att.storage_path:
        return None
    # 归属校验（安全红线）：str 比较兼容 sqlite 字符串存储
    if str(att.project_id) != str(project_id):
        return None
    try:
        from app.core.storage import get_storage

        data = get_storage().get("personal", att.storage_path)
        return data, att.mime_type or "image/png"
    except Exception:
        return None


def inline_share_images(db: Session, content: dict, project_id: uuid.UUID) -> dict:
    """把 Tiptap JSON 里 image 节点的 src 改写为 base64 data URI（游客浏览用）。

    深拷贝后遍历，不修改入参。跨项目 / 取不到字节的图片 src 置空（不泄露、不阻断）。
    """
    import base64
    import copy

    result = copy.deepcopy(content)

    def walk(node: Any) -> None:
        if isinstance(node, dict):
            if node.get("type") == "image":
                attrs = node.setdefault("attrs", {})
                fetched = fetch_image_bytes_for_project(db, attrs.get("src", ""), project_id)
                if fetched:
                    data, mime = fetched
                    b64 = base64.b64encode(data).decode("ascii")
                    attrs["src"] = f"data:{mime};base64,{b64}"
                else:
                    attrs["src"] = ""  # 跨项目/失败：置空，不泄露他人项目附图
            for child in node.get("content", []):
                walk(child)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(result)
    return result


def _get_text(node: dict) -> str:
    if node.get("type") == "text":
        return node.get("text", "")
    return ""
